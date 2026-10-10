# import io, os, sqlite3, datetime
# from flask import Flask, request, jsonify, send_file
# from docx import Document
# from docx.shared import Pt, Cm
# from docx.oxml.ns import qn

# os.makedirs("static", exist_ok=True)

# BASE = os.path.dirname(os.path.abspath(__file__))
# DB_PATH = os.path.join(BASE, "letters.db")

# app = Flask(__name__)

# # ===== DATABASE =====
# def db():
#     con = sqlite3.connect(DB_PATH)
#     con.row_factory = sqlite3.Row
#     return con

# def init_db():
#     try:
#         with db() as con:
#             con.execute("""CREATE TABLE IF NOT EXISTS letters(
#                 id INTEGER PRIMARY KEY AUTOINCREMENT,
#                 client TEXT, sender TEXT, recipient TEXT,
#                 letter_date TEXT, subject TEXT, body TEXT, updated_at TEXT)""")
#         print("✅ Database initialized")
#     except Exception as e:
#         print(f"⚠️ DB Error: {e}")

# init_db()

# # ===== HOME PAGE (EMBEDDED HTML) =====
# HTML_CONTENT = """<!DOCTYPE html>
# <html lang="hi">
# <head>
# <meta charset="utf-8">
# <meta name="viewport" content="width=device-width, initial-scale=1">
# <title>वकील वॉइस</title>
# <link rel="preconnect" href="https://fonts.googleapis.com">
# <link href="https://fonts.googleapis.com/css2?family=Noto+Sans+Devanagari:wght@400;600&family=Noto+Serif+Devanagari:wght@400;600;700&display=swap" rel="stylesheet">
# <style>
# :root{--navy:#0f2742;--gold:#b8893b;--bg:#f4f1ea;--card:#fffdf9;--ink:#1d2733;--mute:#6b7685;--line:#e3ddd0}
# *{box-sizing:border-box}
# body{margin:0;display:flex;min-height:100vh;background:var(--bg);color:var(--ink);font-family:"Noto Sans Devanagari",system-ui,sans-serif;line-height:1.6}
# #side{width:280px;flex-shrink:0;background:linear-gradient(180deg,var(--navy),#0a1b2f);color:#e9eef5;padding:20px 14px;display:flex;flex-direction:column;height:100vh;position:sticky;top:0;overflow:auto}
# .brand{display:flex;gap:10px;align-items:center;margin-bottom:18px}
# .logo{width:40px;height:40px;border-radius:10px;display:grid;place-items:center;font-size:1.4rem;background:linear-gradient(135deg,var(--gold),#d4a95a)}
# .brand h1{margin:0;font-family:"Noto Serif Devanagari",serif;font-size:1.2rem;color:#fff}
# .brand p{margin:0;font-size:.75rem;color:#9fb1c7}
# .profile-btn{width:100%;border:0;border-radius:10px;padding:10px;font:inherit;font-weight:600;cursor:pointer;background:#ffffff1a;color:#fff;margin-bottom:8px;border:1px solid #ffffff22;transition:.15s}
# .profile-btn:hover{background:#ffffff2a}
# .btn-new{width:100%;border:0;border-radius:10px;padding:10px;font:inherit;font-weight:600;cursor:pointer;background:linear-gradient(135deg,var(--gold),#d4a95a);color:var(--navy);margin-bottom:12px}
# .search{width:100%;margin:10px 0 14px;padding:9px;border-radius:10px;border:1px solid #ffffff22;background:#ffffff12;color:#fff;font:inherit;font-size:.9rem}
# #list{list-style:none;margin:0;padding:0;flex:1;overflow:auto}
# #list li{padding:9px;border-radius:8px;cursor:pointer;margin-bottom:3px;border:1px solid transparent;transition:.15s}
# #list li:hover{background:#ffffff14}
# #list b{display:block;font-weight:600;font-size:.9rem;color:#fff}
# #list small{color:#9fb1c7;font-size:.75rem}
# main{flex:1;padding:22px 28px 40px;max-width:950px;margin:0 auto}
# .top{margin-bottom:18px}
# .top h2{margin:0;font-family:"Noto Serif Devanagari",serif;font-size:1.6rem;color:var(--navy)}
# .top p{margin:2px 0 0;color:var(--mute);font-size:.9rem}
# .card{background:var(--card);border:1px solid var(--line);border-radius:14px;padding:18px;margin-bottom:14px}
# .field{margin-bottom:12px;position:relative}
# .field label{display:block;font-size:.8rem;font-weight:600;color:var(--mute);margin-bottom:5px}
# .field-wrapper{position:relative;display:flex;gap:8px}
# .field-wrapper textarea, .field-wrapper input{flex:1}
# input,select,textarea{font:inherit;font-size:.95rem;color:var(--ink);background:#fff;border:1px solid var(--line);border-radius:9px;padding:9px 10px;width:100%;transition:border-color .15s}
# input:focus,select:focus,textarea:focus{outline:none;border-color:var(--gold)}
# textarea{resize:vertical;min-height:150px;font-family:"Noto Serif Devanagari",serif;line-height:1.8}
# .grid{display:grid;grid-template-columns:1fr 1fr;gap:10px;margin-bottom:12px}
# .grid .field{margin:0}
# .grid.full{grid-template-columns:1fr}
# button{font:inherit;cursor:pointer}
# .btn{border-radius:9px;padding:9px 14px;font-weight:500;border:1px solid var(--line);background:#fff;color:var(--navy);transition:.15s}
# .btn:hover{border-color:var(--gold);background:#fffaf0}
# .btn-primary{background:var(--navy);border:0;color:#fff;padding:10px 16px}
# .btn-primary:hover{background:#173a5e}
# .mic-small{border-radius:50%;width:40px;height:40px;display:grid;place-items:center;font-weight:600;background:linear-gradient(135deg,var(--navy),#173a5e);color:#fff;border:0;font-size:1rem;transition:.2s;box-shadow:0 2px 8px rgba(15,39,66,.2);flex-shrink:0}
# .mic-small:hover{transform:translateY(-1px)}
# .mic-small.on{background:linear-gradient(135deg,#c0392b,#e0584a);animation:pulse 1.2s infinite}
# .mic-large{border-radius:50%;width:50px;height:50px;display:grid;place-items:center;font-weight:600;background:linear-gradient(135deg,var(--navy),#173a5e);color:#fff;border:0;font-size:1.3rem;transition:.2s;box-shadow:0 4px 12px rgba(15,39,66,.3)}
# .mic-large:hover{transform:translateY(-2px)}
# .mic-large.on{background:linear-gradient(135deg,#c0392b,#e0584a);animation:pulse 1.2s infinite}
# @keyframes pulse{0%{box-shadow:0 0 0 0 rgba(192,57,43,.5)}70%{box-shadow:0 0 0 12px rgba(192,57,43,0)}100%{box-shadow:0 0 0 0}}
# .actions{display:flex;flex-wrap:wrap;gap:8px;margin-top:12px}
# .interim{color:var(--gold);font-size:.9rem;font-style:italic;min-height:1.4em;margin-top:8px}
# .spell-box{background:#e3f2fd;border-left:4px solid #2196f3;padding:10px 12px;margin-top:10px;border-radius:6px;font-size:.85rem}
# .spell-item{margin-bottom:6px}
# .spell-wrong{color:#c0392b;font-weight:600}
# .spell-right{color:#2c6a47}
# .help{color:var(--mute);font-size:.8rem;margin-top:14px}
# .modal{display:none;position:fixed;top:0;left:0;right:0;bottom:0;background:rgba(0,0,0,.5);z-index:999;align-items:center;justify-content:center}
# .modal.show{display:flex}
# .modal-content{background:var(--card);border-radius:14px;padding:24px;max-width:500px;width:90%;max-height:80vh;overflow:auto}
# .modal-content h3{margin:0 0 16px;color:var(--navy)}
# .modal-close{position:absolute;top:16px;right:16px;border:0;background:none;cursor:pointer;font-size:1.5rem;color:var(--mute)}
# @media(max-width:800px){body{flex-direction:column}#side{width:100%;height:auto;position:static}main{padding:14px}}
# @media print{#side,.actions,.help,.interim,.spell-box,.modal,input,select,.btn,.mic-small,.mic-large{display:none}body{background:#fff}textarea{border:0}}
# </style>
# </head>
# <body>
# <aside id="side">
# <div class="brand"><span class="logo">⚖</span><div><h1>वकील वॉइस</h1><p>आवाज़ से पत्र</p></div></div>
# <button class="profile-btn" onclick="openProfileModal()">👤 प्रोफाइल</button>
# <button class="btn-new" onclick="newLetter()">+ नया पत्र</button>
# <input class="search" id="search" placeholder="🔍 खोजें" oninput="searchLetters()">
# <div style="font-size:.7rem;letter-spacing:.05em;color:#8ea3bb;margin-bottom:8px;text-transform:uppercase">सहेजे पत्र</div>
# <ul id="list"></ul>
# </aside>
# <main>
# <div class="top"><h2>आवाज़ से पत्र लिखें</h2><p>बोलिए — हिंदी में अपने आप टाइप होगा</p></div>
# <div class="card">
# <div class="grid">
# <div class="field"><label>न्यायालय का प्रकार</label>
# <select id="court">
# <option value="">— चुनें —</option>
# <option value="District Court">जिला न्यायालय</option>
# <option value="High Court">उच्च न्यायालय</option>
# <option value="Supreme Court">सर्वोच्च न्यायालय</option>
# </select></div>
# <div class="field"><label>पत्र का प्रकार</label>
# <select id="tpl">
# <option value="">— चुनें —</option>
# <option value="notice">कानूनी नोटिस</option>
# </select></div>
# </div>
# </div>
# <div class="card">
# <div class="grid">
# <div class="field"><label>मुवक्किल का नाम</label><input id="client" placeholder="जैसे: श्री राम कुमार"></div>
# <div class="field"><label>दिनांक</label><input id="date" type="date"></div>
# </div>
# <div class="field">
# <label>अधिवक्ता का नाम व पता</label>
# <div class="field-wrapper">
# <textarea id="from" rows="3" placeholder="आपका नाम, पता, फोन"></textarea>
# <button class="mic-small" id="micFrom" onclick="startFieldVoice('from')">🎤</button>
# </div>
# </div>
# <div class="field">
# <label>प्राप्तकर्ता का नाम व पता</label>
# <div class="field-wrapper">
# <textarea id="to" rows="3" placeholder="किसे भेजना है"></textarea>
# <button class="mic-small" id="micTo" onclick="startFieldVoice('to')">🎤</button>
# </div>
# </div>
# <div class="field">
# <label>विषय</label>
# <div class="field-wrapper">
# <input id="sub" placeholder="पत्र का विषय">
# <button class="mic-small" id="micSub" onclick="startFieldVoice('sub')">🎤</button>
# </div>
# </div>
# </div>
# <div class="card">
# <div style="display:flex;justify-content:space-between;align-items:center;margin-bottom:10px">
# <label style="font-size:.8rem;font-weight:600;color:var(--mute);margin:0">पत्र की सामग्री</label>
# <button class="mic-large" id="micBody" onclick="startFieldVoice('body')">🎤</button>
# </div>
# <textarea id="body" placeholder="यहाँ पत्र लिखें या बोलें..."></textarea>
# <div class="interim" id="interim"></div>
# </div>
# <div class="actions">
# <button class="btn-primary" onclick="saveLetter()">💾 सहेजें</button>
# <button class="btn" onclick="exportWord()">⬇ Word</button>
# <button class="btn" onclick="window.print()">🖨 प्रिंट</button>
# </div>
# </main>

# <script>
# const $=id=>document.getElementById(id);
# let currentId=null, rec=null, on=false, currentField=null;

# $("date").valueAsDate=new Date();

# function newLetter(){currentId=null;$("client").value=$("to").value=$("sub").value=$("body").value="";$("interim").textContent="";}

# async function saveLetter(){
#   const f={id:currentId,client:$("client").value,sender:$("from").value,recipient:$("to").value,letter_date:$("date").value,subject:$("sub").value,body:$("body").value};
#   const r=await fetch("/api/letters",{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify(f)});
#   const j=await r.json();
#   currentId=j.id;
#   refreshList();
#   alert("सहेज लिया ✓");
# }

# async function refreshList(){
#   const r=await fetch("/api/letters");
#   const rows=await r.json();
#   $("list").innerHTML=rows.map(x=>`<li onclick="loadLetter(${x.id})"><b>${x.client||"(बिना नाम)"}</b><small>${x.subject||""}</small></li>`).join("");
# }

# async function loadLetter(id){
#   const r=await fetch("/api/letters/"+id);
#   const d=await r.json();
#   currentId=d.id;
#   $("client").value=d.client;
#   $("from").value=d.sender;
#   $("to").value=d.recipient;
#   $("date").value=d.letter_date;
#   $("sub").value=d.subject;
#   $("body").value=d.body;
# }

# function searchLetters(){
#   const q=$("search").value;
#   if(q)fetch("/api/letters?q="+encodeURIComponent(q)).then(r=>r.json()).then(rows=>$("list").innerHTML=rows.map(x=>`<li onclick="loadLetter(${x.id})"><b>${x.client}</b><small>${x.subject}</small></li>`).join(""));
# }

# async function exportWord(){
#   const f={sender:$("from").value,recipient:$("to").value,letter_date:$("date").value,subject:$("sub").value,body:$("body").value};
#   const r=await fetch("/api/export",{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify(f)});
#   const blob=await r.blob();
#   const url=URL.createObjectURL(blob);
#   const a=document.createElement("a");
#   a.href=url;
#   a.download=($("client").value||"letter")+".docx";
#   a.click();
# }

# const cmds=[[/नया पैराग्राफ|न्यू पैराग्राफ/g,"\\n\\n"],[/नई लाइन|नयी लाइन/g,"\\n"],[/पूर्ण विराम|पूर्णविराम/g,"।"],[/कॉमा|अल्प विराम/g,","]];

# function fix(t){cmds.forEach(([r,v])=>{t=t.replace(r,v)});return t;}

# function addText(t, fieldId=null){
#   const targetField=fieldId?$(fieldId):$("body");
#   if(!targetField)return;
#   const s=targetField.selectionStart??targetField.value.length;
#   const e=targetField.selectionEnd??s;
#   const pre=targetField.value.slice(0,s);
#   const post=targetField.value.slice(e);
#   const sp=(pre&&!/[\\s\\n]$/.test(pre))?" ":"";
#   targetField.value=pre+sp+t+post;
#   const p=(pre+sp+t).length;
#   targetField.setSelectionRange(p,p);
# }

# const SR=window.SpeechRecognition||window.webkitSpeechRecognition;
# if(SR){
#   rec=new SR();rec.lang="hi-IN";rec.continuous=true;rec.interimResults=true;
#   rec.onresult=e=>{let im="";for(let i=e.resultIndex;i<e.results.length;i++){const r=e.results[i];r.isFinal?addText(fix(r[0].transcript.trim()),currentField):im+=r[0].transcript}$("interim").textContent=im?"सुन रहा हूँ: "+im:""};
#   rec.onerror=e=>$("interim").textContent="Error: "+e.error;
#   rec.onend=()=>{if(on)try{rec.start()}catch(_){}};
# }

# function startFieldVoice(fieldId){
#   if(!rec)return alert("Chrome खोलें");
#   currentField=fieldId;
#   if(on){
#     on=false;
#     rec.stop();
#     document.querySelectorAll(".mic-small, .mic-large").forEach(m=>{m.classList.remove("on");m.textContent="🎤"});
#     $("interim").textContent="";
#   }else{
#     on=true;
#     try{rec.start()}catch(_){}
#     document.querySelectorAll(".mic-small, .mic-large").forEach(m=>{m.classList.toggle("on",m.id==="mic"+fieldId.charAt(0).toUpperCase()+fieldId.slice(1))});
#     $("interim").textContent="सुन रहा हूँ...";
#   }
# }

# refreshList();
# </script>
# </body>
# </html>"""

# @app.route("/")
# def home():
#     return HTML_CONTENT

# # ===== API ROUTES =====
# @app.get("/api/letters")
# def list_letters():
#     try:
#         q = request.args.get("q", "").strip()
#         with db() as con:
#             if q:
#                 like = f"%{q}%"
#                 rows = con.execute("""SELECT id,client,subject,letter_date,updated_at FROM letters
#                     WHERE client LIKE ? OR subject LIKE ? ORDER BY id DESC""", (like, like)).fetchall()
#             else:
#                 rows = con.execute("""SELECT id,client,subject,letter_date,updated_at
#                     FROM letters ORDER BY id DESC LIMIT 100""").fetchall()
#         return jsonify([dict(r) for r in rows])
#     except Exception as e:
#         return jsonify(error=str(e)), 500

# @app.get("/api/letters/<int:lid>")
# def get_letter(lid):
#     try:
#         with db() as con:
#             r = con.execute("SELECT * FROM letters WHERE id=?", (lid,)).fetchone()
#         return (jsonify(dict(r)), 200) if r else (jsonify(error="not found"), 404)
#     except Exception as e:
#         return jsonify(error=str(e)), 500

# @app.post("/api/letters")
# def save_letter():
#     try:
#         d = request.get_json(force=True)
#         now = datetime.datetime.now().isoformat(timespec="seconds")
#         vals = (d.get("client", ""), d.get("sender", ""), d.get("recipient", ""),
#                 d.get("letter_date", ""), d.get("subject", ""), d.get("body", ""), now)
#         with db() as con:
#             if d.get("id"):
#                 con.execute("""UPDATE letters SET client=?,sender=?,recipient=?,letter_date=?,
#                     subject=?,body=?,updated_at=? WHERE id=?""", vals + (d["id"],))
#                 lid = d["id"]
#             else:
#                 cur = con.execute("""INSERT INTO letters(client,sender,recipient,letter_date,
#                     subject,body,updated_at) VALUES(?,?,?,?,?,?,?)""", vals)
#                 lid = cur.lastrowid
#         return jsonify(id=lid)
#     except Exception as e:
#         return jsonify(error=str(e)), 500

# @app.delete("/api/letters/<int:lid>")
# def delete_letter(lid):
#     try:
#         with db() as con:
#             con.execute("DELETE FROM letters WHERE id=?", (lid,))
#         return jsonify(ok=True)
#     except Exception as e:
#         return jsonify(error=str(e)), 500

# # ===== WORD EXPORT =====
# def set_hindi_font(run_or_style, name="Mangal"):
#     try:
#         run_or_style.font.name = name
#         el = run_or_style.element.rPr if hasattr(run_or_style.element, "rPr") else None
#         if el is not None:
#             el.rFonts.set(qn("w:cs"), name)
#             el.rFonts.set(qn("w:eastAsia"), name)
#     except:
#         pass

# @app.post("/api/export")
# def export_docx():
#     try:
#         d = request.get_json(force=True)
#         doc = Document()
#         sec = doc.sections[0]
#         sec.top_margin = sec.bottom_margin = Cm(2.5)
#         sec.left_margin = sec.right_margin = Cm(2.5)
#         st = doc.styles["Normal"]
#         st.font.size = Pt(13)
#         set_hindi_font(st)

#         def para(text, bold=False, align=None):
#             p = doc.add_paragraph()
#             r = p.add_run(text)
#             r.bold = bold
#             set_hindi_font(r)
#             p.paragraph_format.space_after = Pt(8)
#             p.paragraph_format.line_spacing = 1.4
#             if align == "right":
#                 p.alignment = 2
#             return p

#         para(d.get("sender", ""), bold=True)
#         if d.get("letter_date"):
#             y, m, dd = d["letter_date"].split("-")
#             para(f"दिनांक: {dd}/{m}/{y}", align="right")
#         if d.get("recipient"):
#             para("सेवा में,\n" + d["recipient"])
#         if d.get("subject"):
#             para("विषय: " + d["subject"], bold=True)
#         para(d.get("body", ""))
#         if d.get("sender"):
#             para("\nभवदीय,\n" + d["sender"].split("\n")[0] + "\nअधिवक्ता")

#         buf = io.BytesIO()
#         doc.save(buf)
#         buf.seek(0)
#         return send_file(buf, as_attachment=True, download_name="letter.docx",
#             mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
#     except Exception as e:
#         return jsonify(error=str(e)), 500

# @app.get("/health")
# def health():
#     return jsonify(status="ok")

# if __name__ == "__main__":
#     port = int(os.environ.get('PORT', 5000))
#     print(f"\n🚀 Vakil Voice running on port {port}\n")
#     app.run(host="0.0.0.0", port=port, debug=False)

import io, os, sqlite3, datetime
from flask import Flask, request, jsonify, send_file
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE, "letters.db")
app = Flask(__name__)


def db():
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    return con


def init_db():
    with db() as con:
        con.execute("""CREATE TABLE IF NOT EXISTS letters(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            client TEXT, sender TEXT, recipient TEXT,
            letter_date TEXT, subject TEXT, body TEXT, updated_at TEXT)""")


init_db()

HTML_CONTENT = r"""<!DOCTYPE html>
<html lang="hi">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>वकील वॉइस</title>
<link href="https://fonts.googleapis.com/css2?family=Noto+Sans+Devanagari:wght@400;600&family=Noto+Serif+Devanagari:wght@400;600;700&display=swap" rel="stylesheet">
<style>
:root{--navy:#0f2742;--gold:#b8893b;--bg:#f4f1ea;--card:#fffdf9;--ink:#1d2733;--mute:#6b7685;--line:#e3ddd0}
*{box-sizing:border-box}
body{margin:0;display:flex;min-height:100vh;background:var(--bg);color:var(--ink);font-family:"Noto Sans Devanagari",system-ui,sans-serif;line-height:1.6}
#side{width:280px;flex-shrink:0;background:linear-gradient(180deg,var(--navy),#0a1b2f);color:#e9eef5;padding:20px 14px;display:flex;flex-direction:column;height:100vh;position:sticky;top:0;overflow:auto}
.brand{display:flex;gap:10px;align-items:center;margin-bottom:18px}
.logo{width:40px;height:40px;border-radius:10px;display:grid;place-items:center;font-size:1.4rem;background:linear-gradient(135deg,var(--gold),#d4a95a)}
.brand h1{margin:0;font-family:"Noto Serif Devanagari",serif;font-size:1.2rem;color:#fff}
.brand p{margin:0;font-size:.75rem;color:#9fb1c7}
.side-btn{width:100%;border-radius:10px;padding:10px;font:inherit;font-weight:600;cursor:pointer;background:#ffffff1a;color:#fff;margin-bottom:8px;border:1px solid #ffffff22}
.btn-new{width:100%;border:0;border-radius:10px;padding:10px;font:inherit;font-weight:600;cursor:pointer;background:linear-gradient(135deg,var(--gold),#d4a95a);color:var(--navy);margin-bottom:12px}
.search{width:100%;margin:6px 0 14px;padding:9px;border-radius:10px;border:1px solid #ffffff22;background:#ffffff12;color:#fff;font:inherit;font-size:.9rem}
#list{list-style:none;margin:0;padding:0;flex:1;overflow:auto}
#list li{padding:9px;border-radius:8px;cursor:pointer;margin-bottom:3px}
#list li:hover{background:#ffffff14}
#list b{display:block;font-weight:600;font-size:.9rem;color:#fff}
#list small{color:#9fb1c7;font-size:.75rem}
main{flex:1;padding:22px 28px 40px;max-width:950px;margin:0 auto;min-width:0}
.top h2{margin:0;font-family:"Noto Serif Devanagari",serif;font-size:1.6rem;color:var(--navy)}
.top p{margin:2px 0 18px;color:var(--mute);font-size:.9rem}
.card{background:var(--card);border:1px solid var(--line);border-radius:14px;padding:18px;margin-bottom:14px}
.field{margin-bottom:12px}
.field label{display:block;font-size:.8rem;font-weight:600;color:var(--mute);margin-bottom:5px}
.wrap{display:flex;gap:8px;align-items:flex-start}
.wrap textarea,.wrap input{flex:1;min-width:0}
input,select,textarea{font:inherit;font-size:.95rem;color:var(--ink);background:#fff;border:1px solid var(--line);border-radius:9px;padding:9px 10px;width:100%}
input:focus,select:focus,textarea:focus{outline:none;border-color:var(--gold)}
textarea{resize:vertical;min-height:90px;font-family:"Noto Serif Devanagari",serif;line-height:1.8}
#body{min-height:260px}
.grid{display:grid;grid-template-columns:1fr 1fr;gap:10px}
button{font:inherit;cursor:pointer}
.btn{border-radius:9px;padding:9px 14px;font-weight:500;border:1px solid var(--line);background:#fff;color:var(--navy)}
.btn-primary{background:var(--navy);border:0;color:#fff;padding:10px 16px;border-radius:9px}
.mic{border-radius:50%;width:42px;height:42px;flex-shrink:0;background:linear-gradient(135deg,var(--navy),#173a5e);color:#fff;border:0;font-size:1.1rem;position:relative}
.mic.big{width:54px;height:54px;font-size:1.4rem}
.mic.on{background:linear-gradient(135deg,#c0392b,#e0584a);animation:pulse 1.2s infinite}
.mic.on::after{content:"ON";position:absolute;top:-18px;right:2px;font-size:.7rem;font-weight:700;color:#c0392b}
@keyframes pulse{0%{box-shadow:0 0 0 0 rgba(192,57,43,.5)}70%{box-shadow:0 0 0 12px rgba(192,57,43,0)}100%{box-shadow:0 0 0 0}}
.status{margin-top:10px;padding:8px 10px;border-radius:6px;background:#fffaf0;border-left:4px solid var(--gold);font-size:.9rem;color:#7a5a1e}
.status.rec{background:#fdecea;border-left-color:#c0392b;color:#c0392b}
.actions{display:flex;flex-wrap:wrap;gap:8px;margin-top:12px}
.help{color:var(--mute);font-size:.8rem;margin-top:14px}
.modal{display:none;position:fixed;inset:0;background:rgba(0,0,0,.5);z-index:99;align-items:center;justify-content:center}
.modal.show{display:flex}
.modal-box{background:var(--card);border-radius:14px;padding:22px;width:92%;max-width:480px;max-height:85vh;overflow:auto}
.modal-box h3{margin:0 0 14px;color:var(--navy)}
@media(max-width:800px){body{flex-direction:column}#side{width:100%;height:auto;position:static}#list{max-height:140px}main{padding:14px}}
@media print{#side,.actions,.help,.status,.mic,.modal,select,.btn{display:none}body{background:#fff}}
</style>
</head>
<body>
<aside id="side">
<div class="brand"><span class="logo">⚖</span><div><h1>वकील वॉइस</h1><p>आवाज़ से पत्र</p></div></div>
<button class="side-btn" onclick="openProfile()">👤 प्रोफाइल</button>
<button class="btn-new" onclick="newLetter()">+ नया पत्र</button>
<input class="search" id="search" placeholder="🔍 खोजें" oninput="searchLetters()">
<ul id="list"></ul>
</aside>

<main>
<div class="top"><h2>आवाज़ से पत्र लिखें</h2><p>बोलिए — हिंदी में अपने आप टाइप होगा</p></div>

<div class="card">
<div class="grid">
<div class="field"><label>न्यायालय का प्रकार</label>
<select id="court">
<option value="">— चुनें —</option>
<option>जिला न्यायालय</option><option>उच्च न्यायालय</option><option>सर्वोच्च न्यायालय</option>
<option>उपभोक्ता न्यायालय</option><option>श्रम न्यायालय</option><option>पारिवारिक न्यायालय</option>
</select></div>
<div class="field"><label>पत्र का प्रकार</label>
<select id="tpl">
<option value="">— चुनें —</option>
<option value="notice_dc">कानूनी नोटिस (वसूली)</option>
<option value="cheque_bounce">चेक बाउंस नोटिस</option>
<option value="defamation">मानहानि नोटिस</option>
<option value="harassment">उत्पीड़न नोटिस</option>
<option value="employment_wrongful">गलत बर्खास्तगी नोटिस</option>
<option value="eviction_notice">बेदखली नोटिस</option>
<option value="divorce_settlement">तलाक समझौता</option>
<option value="will_notice">वसीयत अधिसूचना</option>
<option value="property_partition">संपत्ति विभाजन नोटिस</option>
<option value="debt_recovery">कर्ज वसूली नोटिस</option>
<option value="workplace_harassment">कार्यस्थल उत्पीड़न नोटिस</option>
<option value="gst_notice">GST विवाद नोटिस</option>
<option value="free">खाली पत्र</option>
</select></div>
</div>
</div>

<div class="card">
<div class="grid">
<div class="field"><label>मुवक्किल का नाम</label>
<div class="wrap"><input id="client" placeholder="जैसे: श्री राम कुमार"><button class="mic" data-field="client">🎤</button></div></div>
<div class="field"><label>दिनांक</label><input id="date" type="date"></div>
</div>
<div class="field"><label>अधिवक्ता का नाम व पता</label>
<div class="wrap"><textarea id="from" placeholder="आपका नाम, पता, फोन"></textarea><button class="mic" data-field="from">🎤</button></div></div>
<div class="field"><label>प्राप्तकर्ता का नाम व पता</label>
<div class="wrap"><textarea id="to" placeholder="किसे भेजना है"></textarea><button class="mic" data-field="to">🎤</button></div></div>
<div class="field"><label>विषय</label>
<div class="wrap"><input id="sub" placeholder="पत्र का विषय"><button class="mic" data-field="sub">🎤</button></div></div>
</div>

<div class="card">
<div class="field"><label>पत्र की सामग्री</label>
<div class="wrap"><textarea id="body" placeholder="यहाँ पत्र लिखें या बोलें..."></textarea><button class="mic big" data-field="body">🎤</button></div></div>
<div class="status" id="status">माइक बंद है — 🎤 दबाकर बोलना शुरू करें</div>
</div>

<div class="actions">
<button class="btn-primary" onclick="saveLetter()">💾 सहेजें</button>
<button class="btn" onclick="exportWord()">⬇ Word</button>
<button class="btn" onclick="window.print()">🖨 प्रिंट</button>
</div>
<p class="help"><b>आवाज़ कमांड:</b> "पूर्ण विराम" → । &nbsp;|&nbsp; "कॉमा" → , &nbsp;|&nbsp; "नई लाइन" &nbsp;|&nbsp; "नया पैराग्राफ" &nbsp;|&nbsp; "पैराग्राफ संख्या" → 1. 2. 3. &nbsp;|&nbsp; "माइक बंद करो"</p>
</main>

<div class="modal" id="profileModal"><div class="modal-box">
<h3>वकील की प्रोफाइल</h3>
<div class="field"><label>नाम</label><input id="pName"></div>
<div class="field"><label>बार रजिस्ट्रेशन नं.</label><input id="pBar"></div>
<div class="field"><label>पता</label><textarea id="pAddr" style="min-height:70px"></textarea></div>
<div class="field"><label>फोन</label><input id="pPhone" type="tel"></div>
<div class="actions">
<button class="btn-primary" onclick="saveProfile()">सहेजें</button>
<button class="btn" onclick="closeProfile()">बंद करें</button>
</div></div></div>

<script>
const $ = id => document.getElementById(id);
let currentId = null;

/* ---------- TEMPLATES ---------- */
const TEMPLATES = {
notice_dc:{subject:"कानूनी नोटिस – राशि की वसूली हेतु",body:"माननीय महोदय/महोदया,\n\nमैं अपने मुवक्किल श्री ____ की ओर से आपको यह औपचारिक नोटिस भेज रहा हूँ।\n\n1. यह कि मेरे मुवक्किल ने आपको दिनांक ____ को रुपये ____ की राशि प्रदान की थी।\n\n2. यह कि समझौते के अनुसार दिनांक ____ तक भुगतान किया जाना था, परंतु आपने अभी तक भुगतान नहीं किया है।\n\n3. यह कि बार-बार माँग के बावजूद आप भुगतान करने में विफल रहे हैं।\n\nअतः आपको सूचित किया जाता है कि इस नोटिस की प्राप्ति के 15 दिन के भीतर उक्त राशि का भुगतान करें, अन्यथा आपके विरुद्ध विधिक कार्यवाही की जाएगी।"},
cheque_bounce:{subject:"चेक अनादरित होने पर कानूनी नोटिस (धारा 138, NI Act)",body:"माननीय महोदय/महोदया,\n\n1. यह कि आपने दिनांक ____ को चेक नं. ____ (रुपये ____ का), बैंक ____ का, मेरे मुवक्किल के पक्ष में जारी किया था।\n\n2. यह कि उक्त चेक प्रस्तुत करने पर दिनांक ____ को 'अपर्याप्त निधि' के कारण अनादरित हो गया।\n\n3. यह कि यह कृत्य परक्राम्य लिखत अधिनियम की धारा 138 के अंतर्गत दंडनीय अपराध है।\n\nअतः इस नोटिस की प्राप्ति के 15 दिन के भीतर चेक की राशि का भुगतान करें, अन्यथा आपके विरुद्ध आपराधिक वाद दायर किया जाएगा।"},
defamation:{subject:"मानहानि के संबंध में कानूनी नोटिस",body:"माननीय महोदय/महोदया,\n\n1. यह कि आपने दिनांक ____ को ______ (माध्यम) में मेरे मुवक्किल के विरुद्ध झूठा और अपमानजनक कथन प्रकाशित/प्रसारित किया।\n\n2. यह कि इससे मेरे मुवक्किल की प्रतिष्ठा को गंभीर क्षति पहुँची है।\n\nअतः आपको 7 दिन के भीतर लिखित क्षमा याचना करने तथा उक्त सामग्री हटाने का निर्देश दिया जाता है, अन्यथा विधिक कार्यवाही की जाएगी।"},
harassment:{subject:"उत्पीड़न बंद करने हेतु कानूनी नोटिस",body:"माननीय महोदय/महोदया,\n\n1. यह कि दिनांक ____ से आप मेरे मुवक्किल को लगातार परेशान/उत्पीड़ित कर रहे हैं।\n\n2. यह कि आपके द्वारा किए गए कृत्य निम्नलिखित हैं:\n   - ______\n   - ______\n\n3. यह कि इससे मेरे मुवक्किल को गंभीर मानसिक पीड़ा हुई है।\n\nअतः आपको निर्देश दिया जाता है कि ऐसे सभी कृत्य तुरंत बंद करें, अन्यथा आपके विरुद्ध कानूनी कार्यवाही की जाएगी।"},
employment_wrongful:{subject:"गलत बर्खास्तगी के संबंध में कानूनी नोटिस",body:"माननीय महोदय/महोदया,\n\n1. यह कि मेरे मुवक्किल श्री ____ आपके संस्थान में दिनांक ____ से ____ पद पर कार्यरत थे।\n\n2. यह कि उन्हें दिनांक ____ को बिना किसी उचित कारण, नोटिस या सुनवाई के सेवा से हटा दिया गया।\n\n3. यह कि बकाया वेतन रुपये ____ तथा ग्रेच्युटी रुपये ____ का भुगतान लंबित है।\n\nअतः 30 दिन के भीतर बकाया राशि का भुगतान एवं पुनर्बहाली करें, अन्यथा श्रम न्यायालय में वाद दायर किया जाएगा।"},
eviction_notice:{subject:"परिसर खाली करने हेतु नोटिस",body:"माननीय महोदय/महोदया,\n\nसंपत्ति का विवरण: ______ (पता)\n\n1. यह कि आप उक्त परिसर में किरायेदार के रूप में निवास कर रहे हैं।\n\n2. यह कि दिनांक ____ से रुपये ____ मासिक किराया बकाया है।\n\nअतः आपको 30 दिन का नोटिस दिया जाता है कि बकाया किराया चुकाकर उक्त परिसर खाली करें, अन्यथा बेदखली की कार्यवाही की जाएगी।"},
divorce_settlement:{subject:"परस्पर सहमति से विवाह-विच्छेद का समझौता",body:"यह समझौता दिनांक ____ को श्री ______ (पति) एवं श्रीमती ______ (पत्नी) के मध्य निष्पादित किया जाता है।\n\n1. दोनों पक्ष परस्पर सहमति से विवाह-विच्छेद चाहते हैं।\n\n2. स्थायी भरण-पोषण: रुपये ______\n\n3. बच्चों की अभिरक्षा: ______\n\n4. संपत्ति का बँटवारा: ______\n\n5. दोनों पक्ष भविष्य में एक-दूसरे के विरुद्ध कोई दावा नहीं करेंगे।"},
will_notice:{subject:"वसीयत के संबंध में सार्वजनिक सूचना",body:"सर्वसाधारण को सूचित किया जाता है कि श्री ______ का दिनांक ____ को निधन हो गया।\n\n1. उनके द्वारा दिनांक ____ को वसीयत निष्पादित की गई थी।\n\n2. वसीयत के निष्पादक: ______\n\nयदि किसी को कोई आपत्ति हो तो इस सूचना के प्रकाशन से 30 दिन के भीतर लिखित रूप में अवगत कराएँ।"},
property_partition:{subject:"संपत्ति के विभाजन हेतु नोटिस",body:"माननीय महोदय/महोदया,\n\nसंपत्ति का विवरण: ______ (पता, क्षेत्रफल, खसरा/प्लॉट नं.)\n\n1. यह कि उक्त संपत्ति पक्षकारों की संयुक्त संपत्ति है।\n\n2. यह कि मेरे मुवक्किल का उसमें ____ हिस्सा है।\n\nअतः 30 दिन के भीतर आपसी सहमति से विभाजन करें, अन्यथा न्यायालय में विभाजन वाद दायर किया जाएगा।"},
debt_recovery:{subject:"ऋण की वसूली हेतु कानूनी नोटिस",body:"माननीय महोदय/महोदया,\n\n1. यह कि आपने दिनांक ____ को मेरे मुवक्किल से रुपये ____ ऋण के रूप में लिए थे।\n\n2. यह कि ब्याज दर ____ % वार्षिक तथा भुगतान की तिथि ____ तय हुई थी।\n\n3. यह कि आपने आज तक भुगतान नहीं किया।\n\nअतः 15 दिन के भीतर मूलधन व ब्याज सहित भुगतान करें, अन्यथा वसूली की कार्यवाही की जाएगी।"},
workplace_harassment:{subject:"कार्यस्थल पर उत्पीड़न/भेदभाव के संबंध में नोटिस",body:"माननीय महोदय/महोदया,\n\n1. यह कि मेरे मुवक्किल आपके संस्थान में कार्यरत हैं।\n\n2. यह कि दिनांक ____ से उनके साथ निम्नलिखित व्यवहार किया जा रहा है:\n   - ______\n   - ______\n\n3. यह कि प्रबंधन को शिकायत करने पर भी कोई कार्यवाही नहीं हुई।\n\nअतः 15 दिन के भीतर उचित कार्यवाही करें, अन्यथा विधिक कदम उठाए जाएँगे।"},
gst_notice:{subject:"GST/कर के संबंध में कानूनी नोटिस",body:"माननीय महोदय/महोदया,\n\n1. यह कि दिनांक ____ को मेरे मुवक्किल ने आपसे ______ (सेवा/वस्तु) प्राप्त की।\n\n2. यह कि आपने GST ____ % की दर से वसूला, जबकि सही दर ____ % है।\n\n3. यह कि अधिक वसूली गई राशि रुपये ____ है।\n\nअतः 10 दिन के भीतर उक्त राशि वापस करें, अन्यथा उपभोक्ता न्यायालय में वाद दायर किया जाएगा।"},
free:{subject:"",body:""}
};

$("tpl").addEventListener("change", () => {
  const t = TEMPLATES[$("tpl").value];
  if (!t) return;
  $("sub").value = t.subject;
  $("body").value = t.body;
});

/* ---------- PROFILE ---------- */
function openProfile(){
  $("pName").value = localStorage.getItem("pName") || "";
  $("pBar").value = localStorage.getItem("pBar") || "";
  $("pAddr").value = localStorage.getItem("pAddr") || "";
  $("pPhone").value = localStorage.getItem("pPhone") || "";
  $("profileModal").classList.add("show");
}
function closeProfile(){ $("profileModal").classList.remove("show"); }
function saveProfile(){
  ["pName","pBar","pAddr","pPhone"].forEach(k => localStorage.setItem(k, $(k).value));
  fillFromProfile(true);
  closeProfile();
}
function fillFromProfile(force){
  if (!force && $("from").value) return;
  const parts = [localStorage.getItem("pName"), localStorage.getItem("pBar") ? "बार रजि. नं.: " + localStorage.getItem("pBar") : "", localStorage.getItem("pAddr"), localStorage.getItem("pPhone")].filter(Boolean);
  if (parts.length) $("from").value = parts.join("\n");
}

/* ---------- LETTERS CRUD ---------- */
function newLetter(){
  currentId = null;
  ["client","to","sub","body"].forEach(k => $(k).value = "");
  $("from").value = "";
  $("tpl").value = "";
  $("date").valueAsDate = new Date();
  fillFromProfile(true);
  stopMic("माइक बंद है");
}
async function saveLetter(){
  const f = {id:currentId, client:$("client").value, sender:$("from").value, recipient:$("to").value, letter_date:$("date").value, subject:$("sub").value, body:$("body").value};
  const r = await fetch("/api/letters", {method:"POST", headers:{"Content-Type":"application/json"}, body:JSON.stringify(f)});
  const j = await r.json();
  currentId = j.id;
  refreshList();
  alert("सहेज लिया ✓");
}
function renderList(rows){
  $("list").innerHTML = rows.map(x => `<li onclick="loadLetter(${x.id})"><b>${x.client || "(बिना नाम)"}</b><small>${x.subject || ""}</small></li>`).join("");
}
async function refreshList(){ renderList(await (await fetch("/api/letters")).json()); }
async function searchLetters(){
  const q = $("search").value.trim();
  renderList(await (await fetch("/api/letters?q=" + encodeURIComponent(q))).json());
}
async function loadLetter(id){
  const d = await (await fetch("/api/letters/" + id)).json();
  currentId = d.id;
  $("client").value = d.client || ""; $("from").value = d.sender || ""; $("to").value = d.recipient || "";
  $("date").value = d.letter_date || ""; $("sub").value = d.subject || ""; $("body").value = d.body || "";
}
async function exportWord(){
  const f = {sender:$("from").value, recipient:$("to").value, letter_date:$("date").value, subject:$("sub").value, body:$("body").value};
  const r = await fetch("/api/export", {method:"POST", headers:{"Content-Type":"application/json"}, body:JSON.stringify(f)});
  const blob = await r.blob();
  const a = document.createElement("a");
  a.href = URL.createObjectURL(blob);
  a.download = ($("client").value || "letter") + ".docx";
  a.click();
}

/* ---------- VOICE ---------- */
const IS_MOBILE = /Android|iPhone|iPad|iPod|Mobile/i.test(navigator.userAgent);
const SR = window.SpeechRecognition || window.webkitSpeechRecognition;
let rec = null, recording = false, currentField = null;
let paraNo = 0, lastText = "", lastTime = 0;

const STOP_RE = /(माइक|माईक|माइक्रोफोन)\s*(बंद|बन्द)(\s*(करो|कर|करें|कीजिए))?|mic\s*band(\s*kar(o)?)?|stop\s*recording/i;

function applyCommands(t){
  t = t.replace(/नया\s*पैराग्राफ|न्यू\s*पैराग्राफ/g, "\n\n");
  t = t.replace(/नई\s*लाइन|नयी\s*लाइन|न्यू\s*लाइन/g, "\n");
  t = t.replace(/पूर्ण\s*विराम|पूर्णविराम|फुल\s*स्टॉप/g, "।");
  t = t.replace(/अल्प\s*विराम|अल्पविराम|कॉमा/g, ",");
  t = t.replace(/प्रश्न\s*चिन्ह|प्रश्नचिह्न/g, "?");
  t = t.replace(/पैराग्राफ\s*संख्या/g, () => { paraNo++; return "\n" + paraNo + ". "; });
  return t;
}

function appendText(t){
  const el = $(currentField);
  if (!el || !t) return;
  const cur = el.value;
  const noSpace = !cur || /[\s\n]$/.test(cur) || /^[।,?\n]/.test(t);
  el.value = cur + (noSpace ? "" : " ") + t;
  el.scrollTop = el.scrollHeight;
}

function setStatus(msg, isRec){
  $("status").textContent = msg;
  $("status").classList.toggle("rec", !!isRec);
}

function handleFinal(raw){
  let text = raw.trim();
  if (!text) return;
  // duplicate guard (Android कभी-कभी वही result दोबारा भेजता है)
  const now = Date.now();
  if (text === lastText && now - lastTime < 2500) return;
  lastText = text; lastTime = now;

  let stop = false;
  if (STOP_RE.test(text)) { stop = true; text = text.replace(STOP_RE, "").trim(); }
  if (text) appendText(applyCommands(text));
  if (stop) stopMic("माइक बंद हो गया");
}

if (SR) {
  rec = new SR();
  rec.lang = "hi-IN";
  // Android पर continuous + interim से बार-बार दोहराव होता है, इसलिए mobile पर बंद
  rec.continuous = !IS_MOBILE;
  rec.interimResults = !IS_MOBILE;
  rec.maxAlternatives = 1;

  rec.onresult = (e) => {
    let interim = "";
    for (let i = e.resultIndex; i < e.results.length; i++) {
      const res = e.results[i];
      if (res.isFinal) handleFinal(res[0].transcript);
      else interim += res[0].transcript;
    }
    if (recording) setStatus(interim ? "🎤 सुन रहा हूँ: " + interim : "🎤 माइक चालू है — बोलिए", true);
  };

  rec.onerror = (e) => {
    if (e.error === "no-speech" || e.error === "aborted") return;
    if (e.error === "not-allowed" || e.error === "service-not-allowed") {
      stopMic("❌ माइक की अनुमति नहीं मिली — ब्राउज़र में Allow करें");
    } else {
      setStatus("❌ Error: " + e.error, false);
    }
  };

  rec.onend = () => {
    if (recording) setTimeout(() => { if (recording) { try { rec.start(); } catch (_) {} } }, 250);
  };
}

function markButtons(){
  document.querySelectorAll(".mic").forEach(b => b.classList.toggle("on", recording && b.dataset.field === currentField));
}

function stopMic(msg){
  recording = false;
  currentField = null;
  if (rec) { try { rec.stop(); } catch (_) {} }
  markButtons();
  setStatus(msg || "माइक बंद है", false);
}

function startMic(field){
  if (!rec) { alert("कृपया Google Chrome में खोलें"); return; }
  if (recording && currentField === field) { stopMic("माइक बंद हो गया"); return; }
  if (recording) { try { rec.stop(); } catch (_) {} }
  currentField = field;
  recording = true;
  paraNo = 0; lastText = ""; lastTime = 0;
  markButtons();
  setStatus("🎤 माइक चालू है — बोलिए (बंद करने के लिए कहें \"माइक बंद करो\")", true);
  setTimeout(() => { try { rec.start(); } catch (_) {} }, 300);
}

document.querySelectorAll(".mic").forEach(b => b.addEventListener("click", () => startMic(b.dataset.field)));

/* ---------- INIT ---------- */
$("date").valueAsDate = new Date();
fillFromProfile(false);
refreshList();
</script>
</body>
</html>"""


@app.route("/")
def home():
    return HTML_CONTENT


@app.get("/api/letters")
def list_letters():
    q = request.args.get("q", "").strip()
    with db() as con:
        if q:
            like = f"%{q}%"
            rows = con.execute("""SELECT id,client,subject,letter_date,updated_at FROM letters
                WHERE client LIKE ? OR subject LIKE ? OR body LIKE ? ORDER BY id DESC""", (like, like, like)).fetchall()
        else:
            rows = con.execute("""SELECT id,client,subject,letter_date,updated_at
                FROM letters ORDER BY id DESC LIMIT 100""").fetchall()
    return jsonify([dict(r) for r in rows])


@app.get("/api/letters/<int:lid>")
def get_letter(lid):
    with db() as con:
        r = con.execute("SELECT * FROM letters WHERE id=?", (lid,)).fetchone()
    return (jsonify(dict(r)), 200) if r else (jsonify(error="not found"), 404)


@app.post("/api/letters")
def save_letter():
    d = request.get_json(force=True)
    now = datetime.datetime.now().isoformat(timespec="seconds")
    vals = (d.get("client", ""), d.get("sender", ""), d.get("recipient", ""),
            d.get("letter_date", ""), d.get("subject", ""), d.get("body", ""), now)
    with db() as con:
        if d.get("id"):
            con.execute("""UPDATE letters SET client=?,sender=?,recipient=?,letter_date=?,
                subject=?,body=?,updated_at=? WHERE id=?""", vals + (d["id"],))
            lid = d["id"]
        else:
            cur = con.execute("""INSERT INTO letters(client,sender,recipient,letter_date,
                subject,body,updated_at) VALUES(?,?,?,?,?,?,?)""", vals)
            lid = cur.lastrowid
    return jsonify(id=lid)


@app.delete("/api/letters/<int:lid>")
def delete_letter(lid):
    with db() as con:
        con.execute("DELETE FROM letters WHERE id=?", (lid,))
    return jsonify(ok=True)


def set_hindi_font(obj, name="Mangal"):
    try:
        obj.font.name = name
        rpr = obj.element.rPr
        if rpr is not None:
            rpr.rFonts.set(qn("w:cs"), name)
            rpr.rFonts.set(qn("w:eastAsia"), name)
    except Exception:
        pass


@app.post("/api/export")
def export_docx():
    d = request.get_json(force=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(2.5)
    st = doc.styles["Normal"]
    st.font.size = Pt(13)
    set_hindi_font(st)

    def para(text, bold=False, right=False):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        set_hindi_font(r)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.4
        if right:
            p.alignment = 2
        return p

    if d.get("sender"):
        para(d["sender"], bold=True)
    if d.get("letter_date"):
        y, m, dd = d["letter_date"].split("-")
        para(f"दिनांक: {dd}/{m}/{y}", right=True)
    if d.get("recipient"):
        para("सेवा में,\n" + d["recipient"])
    if d.get("subject"):
        para("विषय: " + d["subject"], bold=True)
    para(d.get("body", ""))
    if d.get("sender"):
        para("\nभवदीय,\n" + d["sender"].split("\n")[0] + "\nअधिवक्ता")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="letter.docx",
                     mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


@app.get("/health")
def health():
    return jsonify(status="ok")


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False)